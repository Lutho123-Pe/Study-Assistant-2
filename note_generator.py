import os
from docx import Document
from fpdf import FPDF
import json
import matplotlib.pyplot as plt
import pandas as pd

class NoteGenerator:
    def __init__(self):
        pass

    def export_txt(self, notes: str, filename: str):
        with open(filename, 'w', encoding='utf-8') as f:
            f.write(notes)

    def export_pdf(self, notes: str, filename: str):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)
        for line in notes.split('\n'):
            pdf.cell(0, 10, txt=line, ln=True)
        pdf.output(filename)

    def export_word(self, notes: str, filename: str):
        doc = Document()
        for line in notes.split('\n'):
            doc.add_paragraph(line)
        doc.save(filename)

    def export_json(self, notes: str, filename: str):
        data = {"notes": notes}
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=4)

    def export_excel(self, notes: str, filename: str):
        # Simple export: each line as a row in Excel
        lines = notes.split('\n')
        df = pd.DataFrame(lines, columns=['Notes'])
        df.to_excel(filename, index=False)

    def export_chart(self, data_dict: dict, filename: str):
        # data_dict: {label: value}
        labels = list(data_dict.keys())
        values = list(data_dict.values())
        plt.figure(figsize=(8,6))
        plt.bar(labels, values)
        plt.title('Visual Chart')
        plt.savefig(filename)
        plt.close()
